import json
import httpx
import pytest

from aitutor.crypto import unb64
from aitutor.education import analyze_code, arithmetic, check_answer
from aitutor.errors import AppError
from aitutor.providers import Provider, validate_config
from aitutor.service import Service

PASSWORD = "parent-password-123"
CHILD_PASSWORD = "student-password-123"
PROFILE = {"region": "Тестовый регион", "city": "Тестовый город", "school": "Тестовая школа", "grade": 7}


class FakeProvider:
    calls = []

    def __init__(self, config):
        self.config = config

    def send_message(self, messages):
        self.calls.append(messages)
        return {"text": "Разберём тему по шагам.", "usage": {"input_tokens": 12, "output_tokens": 8}}

    def close(self):
        pass


@pytest.fixture
def family(tmp_path):
    clock = [100.0]
    app = Service(tmp_path / "app-data", provider_factory=FakeProvider, clock=lambda: clock[0])
    parent = app.register("Родитель", PASSWORD, "PARENT")
    child = app.register("Ученик", CHILD_PASSWORD, "STUDENT", PROFILE)
    app.save_provider({"kind": "ollama", "model": "test-only", "timeout": 10, "revision": 0})
    yield app, parent, child, clock
    app.store.close()


def child_login(app, child):
    app.logout()
    app.login(child["id"], CHILD_PASSWORD)


def test_full_snapshot_encryption_and_reopen(family):
    app, parent, child, _ = family
    raw = app.store.path.read_bytes()
    for marker in ["Тестовая школа", "Тестовый город", PASSWORD, CHILD_PASSWORD, "SQLite format 3"]:
        assert marker.encode() not in raw
    payload = json.loads(raw)
    assert {p["name"] for p in payload["directory"]["profiles"]} == {"Родитель", "Ученик"}
    assert all(set(p) == {"id", "name", "role"} for p in payload["directory"]["profiles"])
    assert not unb64(payload["payload"]["data"]).startswith(b"SQLite format 3")
    app.logout()
    app.login(parent["id"], PASSWORD)
    assert app.snapshot()["students"][0]["id"] == child["id"]


def test_wrong_password_and_recovery_preserves_data(family):
    app, parent, child, _ = family
    app.logout()
    with pytest.raises(AppError):
        app.login(parent["id"], "wrong-password")
    app.recover(parent["id"], parent["recovery_key"], "new-parent-password")
    app.login(parent["id"], "new-parent-password")
    assert app.snapshot()["students"][0]["name"] == "Ученик"


def test_reset_child_and_rotate_recovery(family):
    app, parent, child, _ = family
    app.reset_password(child["id"], "new-student-password")
    replacement = app.rotate_recovery()["recovery_key"]
    app.logout()
    with pytest.raises(AppError):
        app.recover(parent["id"], parent["recovery_key"], PASSWORD)
    app.recover(parent["id"], replacement, PASSWORD)
    app.login(child["id"], "new-student-password")
    assert app.user()["name"] == "Ученик"


def test_native_acl_rejects_other_child_and_parent_commands(family):
    app, _, child, _ = family
    other = app.register("Другой", CHILD_PASSWORD, "STUDENT", PROFILE)
    child_login(app, child)
    with pytest.raises(AppError, match="связи"):
        app.access(other["id"])
    for command, args in [("save_provider", {"config": {}}),
                          ("reset_password", {"student": other["id"], "password": PASSWORD}),
                          ("backup", {"password": PASSWORD})]:
        with pytest.raises(AppError):
            app.dispatch(command, args)
    assert [s["id"] for s in app.snapshot()["students"]] == [child["id"]]


def test_unlinked_child_cannot_call_ai(tmp_path):
    app = Service(tmp_path, provider_factory=FakeProvider)
    app.register("Один", CHILD_PASSWORD, "STUDENT", PROFILE)
    with pytest.raises(AppError):
        app.start_lesson("math-equations")
    assert not app.snapshot()["has_parent"]


def test_second_parent_link_is_one_time(family):
    app, _, child, _ = family
    other = app.register("Второй родитель", PASSWORD, "PARENT")
    token = app.create_invite(child["id"])["token"]
    app.logout()
    app.login(other["id"], PASSWORD)
    assert app.snapshot()["students"] == []
    with pytest.raises(AppError):
        app.access(child["id"], parent_only=True)
    app.link(token)
    assert app.snapshot()["students"][0]["id"] == child["id"]
    with pytest.raises(AppError):
        app.link(token)


def test_provider_revision_and_secret_redaction(family):
    app, _, _, _ = family
    app.save_provider({"kind": "openai", "model": "test", "secret": "secret-marker-XYZ",
                       "revision": 1, "allow_text": True})
    assert "secret-marker-XYZ" not in json.dumps(app.snapshot())
    assert b"secret-marker-XYZ" not in app.store.path.read_bytes()
    with pytest.raises(AppError, match="изменены"):
        app.save_provider({"kind": "ollama", "model": "test", "revision": 1})


def test_cloud_consent_gate_before_http(family):
    app, _, child, _ = family
    app.save_provider({"kind": "openai", "model": "test", "secret": "secret", "revision": 1})
    child_login(app, child)
    before = len(FakeProvider.calls)
    with pytest.raises(AppError, match="не разрешил"):
        app.start_lesson("math-equations")
    assert len(FakeProvider.calls) == before


def test_lesson_idempotency_hint_and_resume(family):
    app, _, child, _ = family
    child_login(app, child)
    lesson = app.start_lesson("math-equations")
    wrong = app.answer(lesson["id"], "6", "operation-wrong-1")
    assert wrong["check"]["status"] == "incorrect"
    assert wrong["session"]["hint_level"] < 5
    app.hint(lesson["id"])
    first = app.answer(lesson["id"], "5", "operation-right-1")
    repeated = app.answer(lesson["id"], "5", "operation-right-1")
    assert repeated == first
    knowledge = app.store.all("knowledge", child["id"])[0]
    assert knowledge["attempts_count"] == 2
    assert knowledge["mastery_score"] < 50  # correct answer had a hint
    with pytest.raises(AppError, match="другим"):
        app.answer(lesson["id"], "999", "operation-right-1")
    app.pause(lesson["id"])
    assert app.resume(lesson["id"])["task"].startswith("Реши уравнение 4x")


def test_unknown_answer_is_not_knowledge_error(family):
    app, _, child, _ = family
    child_login(app, child)
    sid = app.start_lesson("math-equations")["id"]
    result = app.answer(sid, "на фото не видно", "unknown-answer-1")
    assert result["check"]["status"] == "unknown"
    assert app.store.all("knowledge", child["id"]) == []


def test_time_limit_and_grace_are_native(family):
    app, _, child, clock = family
    app.set_controls(child["id"], 5)
    child_login(app, child)
    sid = app.start_lesson("math-powers")["id"]
    # Keep taking actions while monotonic time passes, independent of wall clock.
    for _ in range(3):
        clock[0] += 100
        app.hint(sid)
    assert app.daily(child["id"])["seconds"] == 300
    app.pause(sid)
    with pytest.raises(AppError, match="лимит"):
        app.start_lesson("math-powers")
    app.resume(sid)
    for _ in range(8):
        clock[0] += 100
        app.hint(sid)
    clock[0] += 101
    app.tick()
    assert app.lesson(sid)["status"] == "paused"
    assert app.daily(child["id"])["seconds"] == 1200
    with pytest.raises(AppError):
        app.resume(sid)


def test_completed_lesson_and_history_survive_restart(family):
    app, parent, child, _ = family
    child_login(app, child)
    sid = app.start_lesson("math-equations")["id"]
    for n, answer in enumerate(["5", "4", "-3"]):
        result = app.answer(sid, answer, f"complete-answer-{n}")
    assert result["session"]["status"] == "complete"
    app.logout()
    app.login(parent["id"], PASSWORD)
    assert app.snapshot()["students"][0]["sessions"][0]["correct"] == 3


def test_program_approval_preserves_snapshot(family):
    app, parent, child, _ = family
    child_login(app, child)
    sid = app.start_lesson("math-powers")["id"]
    request = app.request_program({"subject": "algebra", "title": "Тестовый учебник", "authors": "Автор", "publisher": "Издатель"})
    app.logout()
    app.login(parent["id"], PASSWORD)
    app.approve_program(request["id"], True)
    assert app.store.get("session", sid)["program_snapshot"] == {}
    assert app.store.get("program", child["id"] + ":algebra")["title"] == "Тестовый учебник"
    with pytest.raises(AppError, match="рассмотрен"):
        app.approve_program(request["id"], False)


def test_backup_restore_and_corruption_leave_original(family):
    app, parent, child, _ = family
    backup = app.backup("archive-password-123")["file"]
    app.reset_password(child["id"], "changed-child-password")
    before = app.store.path.read_bytes()
    with pytest.raises(AppError):
        app.restore(backup, "wrong", parent["id"], PASSWORD)
    assert app.store.path.read_bytes() == before
    app.restore(backup, "archive-password-123", parent["id"], PASSWORD)
    assert app.actor is None
    app.login(child["id"], CHILD_PASSWORD)
    assert app.user()["name"] == "Ученик"


def test_backup_rejects_traversal_and_student_credentials(family):
    app, _, child, _ = family
    backup = app.backup("archive-password-123")["file"]
    original = app.store.path.read_bytes()
    with pytest.raises(AppError):
        app.restore("../outside.aitbackup")
    with pytest.raises(AppError):
        app.restore(backup, "archive-password-123", child["id"], CHILD_PASSWORD)
    assert app.store.path.read_bytes() == original


def test_atomic_save_failure_rolls_back_database_and_wrappers(family, monkeypatch):
    app, _, child, _ = family
    original = app.store.path.read_bytes()
    old_wrap = json.dumps(app.store.envelope["wraps"][child["id"]])
    def fail(*_):
        raise OSError("synthetic disk failure")
    monkeypatch.setattr("aitutor.storage.atomic_write", fail)
    with pytest.raises(OSError):
        app.reset_password(child["id"], "should-not-be-applied")
    assert json.dumps(app.store.envelope["wraps"][child["id"]]) == old_wrap
    assert app.store.path.read_bytes() == original
    assert not any(e["event"] == "reset_student_password" for e in app.store.all("audit"))


@pytest.mark.parametrize("answer,expected,status", [
    ("1/3+1/6", "1/2", "verified"), ("6 м/с", "6 м/с", "verified"),
    ("6 кг", "6 м/с", "incorrect"), ("5", "6", "incorrect"),
    ("__import__('os').system('calc')", "1", "unknown"), ("истина", "ложь", "incorrect"),
])
def test_bounded_independent_verification(answer, expected, status):
    assert check_answer(answer, expected)["status"] == status


def test_code_is_parsed_never_executed(tmp_path):
    target = tmp_path / "must-not-exist"
    source = f"open({str(target)!r}, 'w').write('bad')"
    assert "не исполнялся" in analyze_code(source)
    assert not target.exists()
    assert "Синтаксическая ошибка" in analyze_code("if x\n  print(x)")
    with pytest.raises(Exception):
        arithmetic("10**999999")


@pytest.mark.parametrize("kind", ["openai", "gigachat", "qwen", "ollama"])
def test_four_provider_contracts_with_mock_transport(kind):
    calls = []
    def handler(request):
        calls.append(request)
        if "oauth" in str(request.url):
            return httpx.Response(200, json={"access_token": "mock-token"})
        body = json.loads(request.content)
        assert body["messages"][0]["content"] == "Тест"
        if kind == "ollama":
            return httpx.Response(200, json={"message": {"content": "Ответ"}, "eval_count": 4})
        return httpx.Response(200, json={"choices": [{"message": {"content": "Ответ"}}],
                                        "usage": {"total_tokens": 4}})
    provider = Provider({"kind": kind, "model": "test", "secret": "sensitive-secret"},
                        transport=httpx.MockTransport(handler))
    try:
        result = provider.send_message([{"role": "user", "content": "Тест"}])
        assert result["text"] == "Ответ"
        assert len(calls) == (2 if kind == "gigachat" else 1)
    finally:
        provider.close()


@pytest.mark.parametrize("status,code", [(401, "authorization"), (429, "quota"), (503, "network")])
def test_provider_errors_never_reflect_secrets(status, code):
    provider = Provider({"kind": "openai", "model": "test", "secret": "private-key"},
        transport=httpx.MockTransport(lambda _: httpx.Response(status, text="private-key")))
    try:
        with pytest.raises(AppError) as error:
            provider.send_message([])
        assert error.value.code == code
        assert "private-key" not in str(error.value)
    finally:
        provider.close()


def test_endpoints_are_restricted():
    with pytest.raises(AppError):
        validate_config({"kind": "ollama", "model": "a", "endpoint": "http://192.168.1.1"})
    with pytest.raises(AppError):
        validate_config({"kind": "openai", "model": "a", "secret": "x", "endpoint": "https://evil.example"})


def test_documents_are_encrypted_and_require_confirmation(family, tmp_path):
    from aitutor.documents import import_file
    app, _, child, _ = family
    child_login(app, child)
    path = tmp_path / "homework.txt"
    path.write_text("sensitive-homework-marker: 2 + 2 = 5", encoding="utf-8")
    item = import_file(app, path)
    assert "original" not in item
    assert b"sensitive-homework-marker" not in app.store.path.read_bytes()
    with pytest.raises(AppError, match="подтвердите"):
        app.check_homework(item["text"], subject="algebra", homework=item["id"])
    app.prepare_homework("algebra", homework=item["id"])
    app.confirm_homework(item["id"], item["text"])
    result = app.check_homework(item["text"], subject="algebra", homework=item["id"])
    assert result["verified"] is False
    assert app.active_lesson is not None  # homework is subject to learning-time accounting


def test_docx_table_extraction_and_file_limits(family, tmp_path):
    from docx import Document
    from aitutor.documents import import_file
    app, _, child, _ = family
    child_login(app, child)
    doc = Document()
    doc.add_paragraph("Задача")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "Время"
    path = tmp_path / "task.docx"
    doc.save(path)
    result = import_file(app, path)
    assert "Время" in result["text"] and "Задача" in result["text"]
    wrong = tmp_path / "not-a-pdf.pdf"
    wrong.write_text("fake PDF")
    with pytest.raises(AppError):
        import_file(app, wrong)


def test_cancelled_provider_result_never_creates_lesson(family):
    app, _, child, _ = family
    class CancelProvider(FakeProvider):
        def send_message(self, messages):
            app.cancel_request()
            return super().send_message(messages)
    app.provider_factory = CancelProvider
    child_login(app, child)
    with pytest.raises(AppError, match="отменён"):
        app.start_lesson("math-equations")
    assert app.store.all("session", child["id"]) == []


def test_homework_cannot_bypass_daily_limit(family):
    app, _, child, _ = family
    child_login(app, child)
    with app.store.transaction():
        daily = app.daily(child["id"])
        daily["seconds"] = 100000
        app.store.put("daily", child["id"], daily, child["id"])
    before = len(FakeProvider.calls)
    work = app.prepare_homework("algebra")
    app.confirm_homework(work["id"], "Условие: 2+2; решение: 5")
    with pytest.raises(AppError, match="лимит"):
        app.check_homework("Условие: 2+2; решение: 5", subject="algebra", homework=work["id"])
    assert len(FakeProvider.calls) == before


def test_diagnosis_adapts_and_moves_between_topics(family):
    app, _, child, _ = family
    child_login(app, child)
    lesson = app.start_lesson("math-equations", mode="subject_diagnose")
    assert lesson["index"] == 1
    wrong = app.answer(lesson["id"], "99", "diagnosis-wrong-1")["session"]
    assert wrong["index"] == 0
    right = app.answer(lesson["id"], "5", "diagnosis-right-2")["session"]
    assert right["index"] == 1
    next_topic = app.answer(lesson["id"], "4", "diagnosis-right-3")["session"]
    assert next_topic["topic"] == "math-powers"
    assert app.store.all("knowledge", child["id"])[0]["topic"] == "math-equations"
