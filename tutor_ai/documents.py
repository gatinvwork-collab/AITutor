"""Bounded native import. Originals are separately encrypted inside the vault."""

import io
import zipfile
from pathlib import Path

from .crypto import seal
from .errors import require
from .service import identity, now

MAX_FILE = 8 * 1024 * 1024
MAX_TEXT = 30_000
EXTRACTOR_VERSION = 2


def normalize_text(value):
    return value.replace("\r\n", "\n").replace("\r", "\n")


def bounded_text(parts):
    """Reject before publication; count Unicode code points, not UTF-8 bytes."""
    result, length = [], 0
    for part in parts:
        part = normalize_text(part)
        length += len(part) + bool(result)
        require(length <= MAX_TEXT,
                "Извлечённый текст превышает 30 000 символов. Разбейте задание на несколько файлов "
                "или частей до 30 000 символов и прикрепите нужную часть.", "text_limit")
        result.append(part)
    return "\n".join(result)


def extract_material(raw, ext):
    """Complete bounded extraction shared by new files and legacy rechecking."""
    text, warning, status = "", "", "complete"
    if ext in (".txt", ".py", ".csv"):
        try:
            decoded = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = raw.decode("utf-8-sig", errors="replace")
            status = "uncertain"
            warning = "Некоторые байты не удалось прочитать как UTF-8: они заменены. Проверьте текст "
            warning += "или сохраните исходный файл в UTF-8 и прикрепите заново."
        text = bounded_text([decoded])
    elif ext == ".pdf":
        require(raw.startswith(b"%PDF-"), "Это не PDF-файл.")
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        require(not reader.is_encrypted and len(reader.pages) <= 30, "PDF: не более 30 страниц, без пароля.")
        text = bounded_text(p.extract_text() or "" for p in reader.pages)
        warning = "Проверьте формулы, таблицы и рисунки по оригиналу. Текстовое извлечение может их потерять."
    elif ext == ".docx":
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            require(sum(i.file_size for i in archive.infolist()) <= 32 * 1024 * 1024
                    and len(archive.infolist()) <= 500, "Документ слишком велик после распаковки.")
        from docx import Document
        document = Document(io.BytesIO(raw))
        parts = (p.text for p in document.paragraphs)
        rows = (" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows)
        from itertools import chain
        text = bounded_text(chain(parts, rows))
        warning = "Проверьте распознанный текст. Изображения и формулы требуют ручного подтверждения."
    elif ext in (".png", ".jpg", ".jpeg", ".webp"):
        from PySide6.QtGui import QImage
        image = QImage.fromData(raw)
        require(not image.isNull() and image.width() * image.height() <= 25_000_000,
                "Изображение повреждено или превышает 25 мегапикселей.")
        status = "manual"
        warning = "Автоматическое OCR ещё не подключено. Перепишите условие и решение, затем подтвердите текст."
    else:
        require(False, "Поддерживаются TXT, PY, CSV, PDF, DOCX, PNG, JPG и WEBP.")
    return {"text": text, "warning": warning, "extraction": {
        "version": EXTRACTOR_VERSION, "status": status, "characters": len(text), "limit": MAX_TEXT,
        "unit": "unicode_code_points", "normalization": "lf", "scope": "text_only"}}


def inspect_homework(service, item):
    """Re-extract legacy material without mutating its original text or results."""
    from .crypto import open_seal
    original = item.get("original")
    require(original or not item.get("filename"), "Оригинал старой работы недоступен. "
            "Прикрепите исходный файл заново перед новой проверкой.", "source_unavailable")
    if not original:
        require(len(item.get("text", "")) <= MAX_TEXT, "Текст старой работы превышает 30 000 символов. "
                "Разбейте задание на части.", "text_limit")
        return None
    metadata = item.get("extraction") or {}
    if metadata.get("version") == EXTRACTOR_VERSION and metadata.get("limit") == MAX_TEXT:
        return None
    raw = open_seal(service.store.attachment_key(), original, item["id"].encode())
    require(len(raw) <= MAX_FILE, "Исходный файл превышает 8 МБ. Разбейте задание на части.")
    return extract_material(raw, Path(item.get("filename", "")).suffix.lower())


def import_file(service, path):
    service.user("STUDENT")
    path = Path(path)
    require(path.is_file() and path.stat().st_size <= MAX_FILE, "Файл должен быть не больше 8 МБ.")
    with path.open("rb") as source:
        raw = source.read(MAX_FILE + 1)
    require(len(raw) <= MAX_FILE, "Файл должен быть не больше 8 МБ.")
    ext = path.suffix.lower()
    extracted = extract_material(raw, ext)
    hid = identity()
    record = {"id": hid, "student": service.actor, "created": now(), "filename": path.name,
              **extracted, "confirmed": False,
              "original": seal(service.store.attachment_key(), raw, hid.encode())}
    with service.store.transaction():
        service.store.put("homework", hid, record, service.actor)
    return {k: v for k, v in record.items() if k != "original"}
